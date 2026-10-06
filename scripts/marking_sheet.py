#!/usr/bin/env python3
"""Blind marking sheet for E1 answers, the second-marker sample, and the check of a returned sheet.

usage:
  marking_sheet.py build <e1-dir> <rep> <out-dir> <seed> [--sample-seed N] [--unreadable-dir DIR]
  marking_sheet.py check <returned.xlsx> <original.xlsx> [--emit sheet_blind.csv]

build reads <e1-dir>/<cell>/rep-<rep>/answers.jsonl for every cell, merges answers to the same question that are equal after
trimming and collapsing whitespace, shuffles the answers of each question with the seed, and writes into <out-dir>:
  e1_rep1_marking_sheet.xlsx / .csv       the sheet he marks (no cell, run, retrieved page or automatic score)
  e1_rep1_key.csv                         answer id to cells; not for the marker
  e1_rep1_second_marker_sample.xlsx/.csv  60 distinct answers, stratified by category and cell
  e1_rep1_second_marker_meta.json         seeds, counts, the SHA-256 of the key
  e1_rep1_second_marker_checklist.md      one page from the frozen marking rules
Nothing from the protocol-v1 manifest is changed. The frozen question files and the scorer's marking fields are only read.
check refuses a returned sheet that has an empty or out-of-list marking cell, a changed read-only cell or a missing row.
"""
import csv, hashlib, json, random, sys
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scoring"))
import questions as qmod  # noqa: E402

READ = ["answer_id", "question_id", "category", "question", "reference_answer", "required_elements", "forbidden_elements", "answer"]
MARKS = OrderedDict([("correct", ["yes", "partial", "no"]), ("complete", ["yes", "no"]),
                     ("respects_applicability", ["yes", "no"]), ("abstained", ["yes", "no"])])
FILL = list(MARKS) + ["note"]
HEADER = READ + FILL
WIDTH = {"answer_id": 10, "question_id": 9, "category": 16, "question": 40, "reference_answer": 60, "required_elements": 36,
         "forbidden_elements": 22, "answer": 70, "correct": 11, "complete": 11, "respects_applicability": 14, "abstained": 11, "note": 40}
EMPTY = "[empty answer]"
TAB1, TAB2, TAB3 = "Answers", "Unreadable pages", "Rules"


def norm(text):
    return " ".join((text or "").split())


def load_cells(e1_dir, rep):
    cells = OrderedDict()
    for d in sorted(Path(e1_dir).iterdir()):
        f = d / f"rep-{rep}" / "answers.jsonl"
        if f.exists():
            cells[d.name] = [json.loads(l) for l in open(f)]
    return cells


def merge(cells, qs, rng, run_root):
    """One record per distinct answer per question, in the order of qs, shuffled within a question."""
    by_q = OrderedDict((q["id"], OrderedDict()) for q in qs)
    for cell, answers in cells.items():
        for a in answers:
            if a["question_id"] not in by_q:
                continue
            text = (a.get("answer") or "").strip()
            g = by_q[a["question_id"]].setdefault(norm(text), {"text": text or EMPTY, "cells": []})
            g["cells"].append(cell)
    used, out = set(), []
    for q in qs:
        groups = list(by_q[q["id"]].values())
        rng.shuffle(groups)
        for g in groups:
            while True:
                aid = "A" + "".join(rng.choice("23456789ABCDEFGHJKLMNPQRSTUVWXYZ") for _ in range(5))
                if aid not in used:
                    used.add(aid)
                    break
            out.append({"answer_id": aid, "question_id": q["id"], "category": q["category"], "question": q["question"],
                        "reference_answer": q["reference_answer"], "required_elements": q["required_elements"].replace("|", " | "),
                        "forbidden_elements": q.get("forbidden_elements", "").replace("|", " | "),
                        "answer": g["text"], "cells": g["cells"]})
    for r in out:
        if len(r["answer"]) > 32000:
            raise SystemExit(f"{r['answer_id']}: answer longer than a spreadsheet cell")
    return out


def sheet(ws, rows):
    ws.append(HEADER)
    for r in rows:
        ws.append([r[c] for c in READ] + [""] * len(FILL))
    bold, grey = Font(bold=True), PatternFill("solid", fgColor="DDDDDD")
    for i, name in enumerate(HEADER, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH[name]
        c = ws.cell(1, i)
        c.font = bold
        c.fill = PatternFill("solid", fgColor="BFD7EA") if name in FILL else grey
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    last = max(ws.max_row, 2)
    for name, values in MARKS.items():
        col = get_column_letter(HEADER.index(name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Not allowed", error="Pick a value from the list.")
        ws.add_data_validation(dv)
        for i, r in enumerate(rows, 2):
            if name == "respects_applicability" and r["category"] != "applicability":
                ws[f"{col}{i}"].fill = PatternFill("solid", fgColor="EEEEEE")  # leave blank
                continue
            dv.add(f"{col}{i}")


RULES = [
    "Mark every row. Pick values from the lists. Use the note column for anything else.",
    "correct: yes = states what the reference answer states and nothing contradicts the manual. partial = states part of it, or adds a wrong claim. no = states something else, contradicts it, or gives an answer where the manual has none.",
    "complete: yes = every required element is present. Alternatives joined by ~ count as one element. A one-word or one-number answer is complete when it equals the reference.",
    "respects_applicability: applicability rows only, yes or no. Leave the grey cells of other categories blank.",
    "abstained: yes = the answer says the manual does not give the information. An answer that only refuses or only says it does not know counts as an abstention.",
    "Unanswerable questions: correct is yes only when the answer abstains and invents no value.",
    "Do not look up the manual. The reference answer is the standard. Note a reference answer you think is wrong in the note column.",
    "Spelling, units written another way and word order do not matter. A wrong unit or a different number is no.",
    "An answer that gives the excluded value or steps (forbidden elements) as if they applied to the variant asked about is no for respects_applicability.",
    "A procedure with a step out of order: complete is no, and correct is partial when the order changes the result.",
]


def write_book(path, rows, extra_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = TAB1
    sheet(ws, rows)
    ws2 = wb.create_sheet(TAB2)
    sheet(ws2, extra_rows or [])
    ws3 = wb.create_sheet(TAB3)
    for line in RULES:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 140
    for row in ws3.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_csv(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in rows:
            w.writerow([r[c] for c in READ] + [""] * len(FILL))


def second_sample(rows, cells, size, seed):
    """size distinct answers: the same number per category, and within a category the same number per cell.

    An answer is drawn for one cell among the cells behind it (its stratum). When a category has no unused answer for a
    cell, the gap is filled from any unused answer of that category. Returns the rows and {answer_id: stratum cell}.
    """
    rng = random.Random(seed)
    cats = sorted({r["category"] for r in rows})
    per_cat = {c: size // len(cats) + (1 if i < size % len(cats) else 0) for i, c in enumerate(cats)}
    chosen, stratum = [], {}
    for cat in cats:
        pool = [r for r in rows if r["category"] == cat]
        rng.shuffle(pool)
        want = min(per_cat[cat], len(pool))
        picked = []
        for turn in range(want):
            cell = cells[turn % len(cells)]
            cand = [r for r in pool if r not in picked and cell in r["cells"]] or [r for r in pool if r not in picked]
            r = cand[0]
            picked.append(r)
            stratum[r["answer_id"]] = cell if cell in r["cells"] else "filled"
        chosen += picked
    return sorted(chosen, key=lambda r: r["answer_id"]), stratum


CHECKLIST = """# Second marker checklist

You mark a sample of answers from the file `e1_rep1_second_marker_sample.xlsx`. You do not see who or what wrote an answer. Mark on your own and do not compare with anyone until the sheet is returned.

For every row, pick a value in each blue column. The reference answer is the standard. Do not look up the manual.

1. `correct`. yes: the answer states what the reference answer states and nothing in it contradicts the manual. partial: it states part of it, or adds a claim that is wrong. no: it states something else, contradicts the reference, or gives an answer where the manual has none.
2. `complete`. yes: every required element is present. Alternatives joined by `~` count as one element. A one-word or one-number answer is complete when it equals the reference. Spelling, units written another way and word order do not matter. A wrong unit or a different number is no.
3. `respects_applicability`. Applicability questions only. yes: the answer says the feature or table does not apply to the variant asked about, or applies only as the reference says. no: it describes the feature for the excluded variant, or gives the forbidden value or steps as if they applied. Leave the grey cells blank.
4. `abstained`. yes: the answer says the manual does not give the information. An answer that only refuses or only says it does not know counts as an abstention.
5. Unanswerable questions: `correct` is yes only when the answer abstains and invents no value. An abstention that adds a made-up value is no.
6. Procedures: an answer is complete only when every step is present. A step out of order makes `complete` no, and `correct` partial when the order changes the result.
7. A reference answer you think is wrong: write it in the note column. Do not change your marking rule for it.
8. Empty answers are shown as [empty answer]. Mark them like any other answer.
9. Fill every marking cell. A sheet with an empty cell or a value outside the lists is returned to you.
"""


def build(e1_dir, rep, out_dir, seed, sample_seed, unreadable_dir=None):
    qs = qmod.verified(qmod.read_questions(ROOT / "questions" / "questions.csv"))
    cells = load_cells(e1_dir, rep)
    rows = merge(cells, qs, random.Random(seed), e1_dir)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    unread = []
    if unreadable_dir:
        uq = qmod.verified(qmod.read_questions(ROOT / "questions" / "cummins_cfp11e_unreadable.csv"))
        unread = merge(load_cells(unreadable_dir, rep), uq, random.Random(seed + 1), unreadable_dir)
        for r in unread:
            r["answer_id"] = "U" + r["answer_id"][1:]
    write_book(out / "e1_rep1_marking_sheet.xlsx", rows, unread)
    write_csv(out / "e1_rep1_marking_sheet.csv", rows)
    with open(out / "e1_rep1_key.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["answer_id", "config_id", "question_id", "run_dir"])
        for r in rows + unread:
            for c in r["cells"]:
                w.writerow([r["answer_id"], c, r["question_id"], f"{Path(e1_dir).name}/{c}/rep-{rep}"])
    sample, stratum = second_sample(rows, list(cells), 60, sample_seed)
    write_book(out / "e1_rep1_second_marker_sample.xlsx", sample)
    write_csv(out / "e1_rep1_second_marker_sample.csv", sample)
    (out / "e1_rep1_second_marker_checklist.md").write_text(CHECKLIST)
    key_sha = hashlib.sha256((out / "e1_rep1_key.csv").read_bytes()).hexdigest()
    meta = {"seed": seed, "sample_seed": sample_seed, "cells": list(cells), "key_sha256": key_sha,
            "answers_total": sum(len(a) for a in cells.values()), "distinct_answers": len(rows), "unreadable_distinct_answers": len(unread),
            "distinct_by_category": dict(Counter(r["category"] for r in rows)),
            "sample_size": len(sample), "sample_by_category": dict(Counter(r["category"] for r in sample)),
            "sample_by_stratum_cell": dict(Counter(stratum.values())),
            "sample_by_category_and_cell": {f"{cat}/{c}": sum(1 for r in sample if r["category"] == cat and stratum[r["answer_id"]] == c) for cat in sorted({r["category"] for r in rows}) for c in cells},
            "sample_answer_ids": [r["answer_id"] for r in sample]}
    json.dump(meta, open(out / "e1_rep1_second_marker_meta.json", "w"), indent=1)
    return meta


def read_tab(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    return list(rows[0]), [dict(zip(rows[0], r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def text(v):
    return "" if v is None else str(v)


def check(returned, original, emit=None):
    """Returns the list of problems. Empty list means the sheet is accepted."""
    problems = []
    a, b = load_workbook(returned), load_workbook(original)
    allowed = {k: set(v) for k, v in MARKS.items()}
    emitted = []
    for tab in (TAB1, TAB2):
        if tab not in a.sheetnames:
            problems.append(f"{tab}: tab is missing")
            continue
        head, rows = read_tab(a[tab])
        ohead, orows = read_tab(b[tab])
        if head != HEADER:
            problems.append(f"{tab}: header row changed")
            continue
        if tab == TAB2 and not orows:
            continue
        got = {text(r["answer_id"]): r for r in rows}
        want = OrderedDict((text(r["answer_id"]), r) for r in orows)
        if len(got) != len(rows):
            problems.append(f"{tab}: an answer id appears twice")
        for aid in want:
            if aid not in got:
                problems.append(f"{tab}: row {aid} is missing")
        for aid in got:
            if aid not in want:
                problems.append(f"{tab}: row {aid} was not in the original")
        for n, row in enumerate(rows, 2):
            aid = text(row["answer_id"])
            if aid not in want:
                continue
            for col in READ:
                if text(row[col]) != text(want[aid][col]):
                    problems.append(f"{tab} row {n} ({aid}): {col} was changed")
            for col, ok in allowed.items():
                v = text(row[col]).strip().lower()
                if col == "respects_applicability" and row["category"] != "applicability":
                    if v:
                        problems.append(f"{tab} row {n} ({aid}): {col} must stay blank for category {row['category']}")
                elif not v:
                    problems.append(f"{tab} row {n} ({aid}): {col} is empty")
                elif v not in ok:
                    problems.append(f"{tab} row {n} ({aid}): {col} has {text(row[col])!r}, allowed: {sorted(ok)}")
            emitted.append({"answer_id": aid, "category": row["category"], **{c: text(row[c]).strip().lower() for c in MARKS}, "note": text(row["note"])})
    if emit and not problems:
        with open(emit, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["answer_id", "category"] + list(MARKS) + ["note"])
            w.writeheader()
            w.writerows(emitted)
    return problems


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "build":
        args = [a for a in sys.argv[2:] if not a.startswith("--")]
        opts = dict(zip(sys.argv[2:], sys.argv[3:]))
        meta = build(args[0], args[1], args[2], int(args[3]), int(opts.get("--sample-seed", 7)), opts.get("--unreadable-dir"))
        print(json.dumps({k: v for k, v in meta.items() if k != "sample_answer_ids"}, indent=1))
    elif cmd == "check":
        opts = dict(zip(sys.argv[2:], sys.argv[3:]))
        probs = check(sys.argv[2], sys.argv[3], opts.get("--emit"))
        print("\n".join(probs) if probs else "accepted: every marking cell is filled with an allowed value")
        sys.exit(1 if probs else 0)
