#!/usr/bin/env python3
"""Build the verification pack for the person who checks the reference answers.

usage: make_verification_pack.py <draft.csv> <out-dir>
The draft is the question file plus a `passage` column with the supporting text of the manual. The pack has one row per
question with the reference answer, the pages, the passage, and empty columns for a verdict and a correction. Because the
passages are text of the manual, the draft and the pack stay out of the public repository.
"""
import csv, os, sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = ["id", "category", "question", "reference_answer", "source_pages", "exclusion_pages", "required_elements", "passage"]
VERDICTS = "correct,needs correction,question unclear,reject"


def main(draft, out_dir):
    rows = list(csv.DictReader(open(draft, newline="")))
    missing = [c for c in COLUMNS if c not in rows[0]]
    if missing:
        raise SystemExit(f"missing columns: {missing}")
    os.makedirs(out_dir, exist_ok=True)
    header = COLUMNS + ["verdict", "correction", "verified_by", "verified_on"]
    with open(f"{out_dir}/verification_pack.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({**r, "verdict": "", "correction": "", "verified_by": "", "verified_on": ""})
    wb = Workbook()
    ws = wb.active
    ws.title = "questions"
    ws.append(header)
    for r in rows:
        ws.append([r[c] for c in COLUMNS] + ["", "", "", ""])
    widths = {"id": 9, "category": 18, "question": 44, "reference_answer": 56, "source_pages": 12, "exclusion_pages": 12,
              "required_elements": 36, "passage": 80, "verdict": 18, "correction": 44, "verified_by": 14, "verified_on": 12}
    for i, name in enumerate(header, 1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = widths[name]
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="DDDDDD")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1=f'"{VERDICTS}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{ws.cell(row=2, column=header.index('verdict') + 1).coordinate}:{ws.cell(row=len(rows) + 1, column=header.index('verdict') + 1).coordinate}")
    ws.freeze_panes = "D2"
    wb.save(f"{out_dir}/verification_pack.xlsx")
    return len(rows)


if __name__ == "__main__":
    print(main(*sys.argv[1:3]), "rows")
