#!/usr/bin/env python3
"""Merge the maintainer's verification of a pack into the question files.

usage: merge_verification.py <delivered.csv> <returned.xlsx|csv> <main.csv> <unreadable.csv> --by NAME --on DATE
`delivered` is the pack as it was handed over. A cell the maintainer changed in the returned file is a correction and
replaces the value in the question file. A cell he left alone keeps the value in the question file, which may have been
corrected since delivery. The verdict column: ok marks the row verified; fix marks it verified only when the row has
cell changes or a non-empty correction (see parse_correction); drop removes the row; an empty verdict leaves it
unverified. Prints a report and writes both question files. Exit status 1 when anything is unresolved.
"""
import csv, re, sys
from pathlib import Path

PAGE_MARKER = re.compile(r"PDF pages?\s+(\d+(?:\s*[-,;]\s*\d+)*)\s*\.")
NOTE_STARTS = ("The original", "Remove the", "The question remains", "Note:")
EDITABLE = ["question", "reference_answer", "source_pages", "exclusion_pages", "required_elements", "forbidden_elements"]


def parse_correction(text):
    """A correction cell: an optional `Question:` line, then `Answer:` up to the `PDF page(s) N.` marker.

    The marker gives the source pages. What follows the marker is the verifier's note to us and is not part of the answer.
    Without a marker (unanswerable rows) the answer ends before the first sentence that starts like a note.
    Returns {"question", "answer", "pages", "notes"}, with None for what the cell does not say.
    """
    out = {"question": None, "answer": None, "pages": None, "notes": ""}
    body = (text or "").strip()
    m = re.match(r"Question:\s*(.*?)\s*\n\s*Answer:\s*(.*)", body, re.S)
    if m:
        out["question"], body = m.group(1).strip(), "Answer: " + m.group(2)
    if not body.startswith("Answer:"):
        out["notes"] = body
        return out
    body = body[len("Answer:"):].strip()
    mark = PAGE_MARKER.search(body)
    if mark:
        out["answer"], out["notes"] = body[:mark.start()].strip(), body[mark.end():].strip()
        nums = [int(n) for n in re.findall(r"\d+", mark.group(1))]
        pages = set(range(nums[0], nums[1] + 1)) if "-" in mark.group(1) and len(nums) == 2 else set(nums)
        out["pages"] = ";".join(str(n) for n in sorted(pages))
    else:
        cut = min([i for i in (body.find(" " + n) for n in NOTE_STARTS) if i >= 0] or [len(body)])
        out["answer"], out["notes"] = body[:cut].strip(), body[cut:].strip()
    return out


def read_any(path):
    if str(path).endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h) for h in rows[0]]
        return [{h: ("" if v is None else str(v)) for h, v in zip(head, r)} for r in rows[1:] if r and r[0]]
    return list(csv.DictReader(open(path, newline="")))


def write(path, rows, fields):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def merge(delivered, returned, main, unreadable, by, on):
    sent = {r["id"]: r for r in delivered}
    files = {"main": main, "unreadable": unreadable}
    report = {"verified": [], "corrected": {}, "dropped": [], "unverified": [], "unresolved": [], "unknown": []}
    for r in returned:
        qid = r["id"]
        if qid not in sent:
            report["unknown"].append(qid)
            continue
        target = {q["id"]: q for q in files.get(sent[qid].get("set", "main"), main)}
        q = target.get(qid)
        if q is None:
            report["unknown"].append(qid)
            continue
        changed = [c for c in EDITABLE if (r.get(c) or "").strip() != (sent[qid].get(c) or "").strip()]
        for c in changed:
            q[c] = r[c].strip()
        verdict = (r.get("verdict") or "").strip().lower()
        if verdict == "fix" and (r.get("correction") or "").strip():
            fix = parse_correction(r["correction"])
            for field, value in (("question", fix["question"]), ("reference_answer", fix["answer"])):
                if value and q[field] != value:
                    q[field] = value
                    changed.append(field)
            if fix["pages"] and q["source_pages"] != fix["pages"] and q["category"] != "unanswerable":
                q["source_pages"] = fix["pages"]
                changed.append("source_pages")
            if fix["notes"]:
                report.setdefault("notes", {})[qid] = fix["notes"]
            changed = changed or ["correction"]
        if changed:
            report["corrected"][qid] = changed
        if verdict == "drop":
            q["_drop"] = True
            report["dropped"].append(qid)
        elif verdict == "ok" or (verdict == "fix" and changed):
            q["verified_by"], q["verified_on"] = (r.get("verified_by") or by), (r.get("verified_on") or on).split(" ")[0]
            report["verified"].append(qid)
        elif verdict == "fix":
            report["unresolved"].append(f"{qid}: verdict fix without a changed cell; correction note: {(r.get('correction') or '').strip()[:120]}")
        else:
            report["unverified"].append(qid)
    return report


def main(argv):
    delivered, returned, main_csv, unread_csv = argv[:4]
    by, on = argv[argv.index("--by") + 1], argv[argv.index("--on") + 1]
    main_rows, unread_rows = read_any(main_csv), read_any(unread_csv)
    fields = list(main_rows[0])
    report = merge(read_any(delivered), read_any(returned), main_rows, unread_rows, by, on)
    for path, rows in ((main_csv, main_rows), (unread_csv, unread_rows)):
        write(path, [{k: v for k, v in q.items() if k != "_drop"} for q in rows if not q.get("_drop")], fields)
    for k, v in report.items():
        print(f"{k}: {len(v)}", dict(v) if isinstance(v, dict) and k != "notes" else (list(v) if isinstance(v, dict) else v))
    return 1 if report["unresolved"] or report["unknown"] else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
