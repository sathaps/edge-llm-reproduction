import csv, json
import marks, retrieval_score, blind_export, questions as qmod


def question(qid, cat, src="2", exc="", elems="a|b", by="x"):
    return {"id": qid, "category": cat, "question": f"question {qid}", "reference_answer": f"reference {qid}", "source_pages": src,
            "exclusion_pages": exc, "required_elements": elems, "forbidden_elements": "z" if cat == "applicability" else "", "verified_by": by, "verified_on": "2026-01-01" if by else ""}


def full_set(n=8):
    cats = list(qmod.MINIMUM)
    rows = []
    for c in cats:
        for i in range(n):
            rows.append(question(f"{c}-{i}", c, src="" if c == "unanswerable" else "3", exc="4" if c == "applicability" else ""))
    return rows


def test_parse_pages():
    assert qmod.parse_pages("12;14-15") == {12, 14, 15}
    assert qmod.parse_pages("") == set()
    assert qmod.parse_pages("3, 5") == {3, 5}


def test_validate_accepts_a_full_set():
    assert qmod.validate(full_set()) == ([], [])


def test_validate_reports_each_problem():
    rows = full_set(5)
    rows[0]["category"] = "other"
    rows[1]["source_pages"] = ""
    rows[2]["id"] = rows[3]["id"]
    errs, warns = qmod.validate(rows)
    joined = " ".join(errs)
    assert "unknown category" in joined and "no source_pages" in joined and "duplicate id" in joined
    assert any("target 8" in w for w in warns)


def test_validate_flags_missing_applicability_exclusion_and_small_categories():
    rows = [r for r in full_set(8) if r["category"] != "table_lookup"]
    for r in rows:
        if r["category"] == "applicability":
            r["exclusion_pages"] = ""
    errs, _ = qmod.validate(rows)
    assert any("no exclusion_pages" in e for e in errs) and any("table_lookup: 0 questions" in e for e in errs)


def retrieval(qid, *page_sets):
    return {"question_id": qid, "retrieved": [{"chunk_id": f"c{i}", "pages": p, "score": 1.0} for i, p in enumerate(page_sets)]}


def test_retrieval_hits_exclusions_and_unverified_are_handled():
    qs = [question("s1", "self_contained", src="3;4"), question("a1", "applicability", src="5", exc="9"),
          question("u1", "unanswerable", src=""), question("s2", "self_contained", by="")]
    rows = retrieval_score.score_run(qs, [retrieval("s1", [1], [3]), retrieval("a1", [5, 6]), retrieval("u1"), retrieval("s2", [3])], "cfg")
    by = {r["question_id"]: r for r in rows}
    assert set(by) == {"s1", "a1", "u1"}
    assert by["s1"]["any_reference_page"] and not by["s1"]["all_reference_pages"] and by["s1"]["first_hit_rank"] == 2
    assert by["a1"]["any_reference_page"] and by["a1"]["exclusion_retrieved"] is False
    assert by["u1"]["any_reference_page"] is None and by["u1"]["n_retrieved"] == 0
    s = retrieval_score.summarise(rows)
    assert s["self_contained"]["N"] == 1 and s["self_contained"]["any"] == 1 and s["self_contained"]["all"] == 0
    assert s["applicability"]["exclusion"] == 0 and s["applicability"]["exclusion_N"] == 1
    assert s["unanswerable"]["returned_nothing"] == 1


def make_run(tmp_path, name, answers):
    d = tmp_path / name
    d.mkdir()
    with open(d / "answers.jsonl", "w") as f:
        for qid, a in answers.items():
            f.write(json.dumps({"question_id": qid, "answer": a}) + "\n")
    with open(d / "retrieval.jsonl", "w") as f:
        for qid in answers:
            f.write(json.dumps({"question_id": qid, "retrieved": [{"chunk_id": "c0", "pages": [3], "score": 1.0}]}) + "\n")
    with open(d / "chunks.jsonl", "w") as f:
        f.write(json.dumps({"chunk_id": "c0", "pages": [3], "text": f"text of {name}"}) + "\n")
    return d


def test_blind_export_shuffles_hides_configuration_and_round_trips(tmp_path):
    qs = [question(f"q{i}", "self_contained", elems="alpha|beta") for i in range(6)]
    answers = {f"q{i}": f"alpha and beta {i}" if i % 2 else "I do not know" for i in range(6)}
    runs = {"A-llama3": make_run(tmp_path, "ra", answers), "B-llama3": make_run(tmp_path, "rb", answers)}
    out = tmp_path / "marking"
    key = blind_export.export(runs, qs, out, seed=7)
    key2 = blind_export.export(runs, qs, tmp_path / "again", seed=7)
    assert key == key2
    assert json.load(open(out / "export.json")) == {"seed": 7, "answers_shown": 6, "answers_total": 12}

    blind = marks.read_csv(out / "sheet_blind.csv")
    assert len(blind) == 6 and len({r["answer_id"] for r in blind}) == 6  # identical answers are shown once
    text = open(out / "sheet_blind.csv").read() + open(out / "sheet_support.csv").read() + open(out / "rubric_marks.csv").read()
    assert "A-llama3" not in text and "B-llama3" not in text and "ra" not in {r["answer_id"] for r in blind}
    configs_in_order = [r["config_id"] for r in marks.read_csv(out / "key.csv")]
    assert configs_in_order != sorted(configs_in_order)

    auto = {r["answer_id"]: r for r in marks.read_csv(out / "rubric_marks.csv")}
    for r in blind:
        a = auto[r["answer_id"]]
        assert (a["abstained"] == "yes") == ("do not know" in r["answer"])
        assert (a["complete"] == "yes") == ("alpha and beta" in r["answer"])

    for r in blind:
        r["correct"] = "yes" if "alpha" in r["answer"] else "no"
        r["complete"] = "yes" if "alpha" in r["answer"] else "no"
    for name, rows in (("sheet_blind.csv", blind),):
        with open(out / name, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    marked = marks.join_marks(marks.read_csv(out / "sheet_blind.csv"), marks.read_csv(out / "sheet_support.csv"), marks.read_csv(out / "key.csv"))
    assert len(marked) == 12  # one mark per answer shown, applied to every configuration behind it
    table = marks.counts(marked)
    assert table[("A-llama3", "self_contained")] == {"N": 6, "yes": 3, "partial": 0}
    assert table[("B-llama3", "self_contained")]["N"] == 6
    configs, rows = marks.outcomes_by_question(marked)
    assert configs == ["A-llama3", "B-llama3"] and len(rows) == 6
    assert all(r["A-llama3"] == r["B-llama3"] for r in rows)


def test_unmarked_rows_are_left_out(tmp_path):
    qs = [question("q0", "self_contained")]
    runs = {"A": make_run(tmp_path, "ra", {"q0": "alpha"})}
    out = tmp_path / "m"
    blind_export.export(runs, qs, out, seed=1)
    marked = marks.join_marks(marks.read_csv(out / "sheet_blind.csv"), marks.read_csv(out / "sheet_support.csv"), marks.read_csv(out / "key.csv"))
    assert marked == []


def test_export_skips_unverified_questions(tmp_path):
    qs = [question("q0", "self_contained", by=""), question("q1", "self_contained")]
    runs = {"A": make_run(tmp_path, "ra", {"q0": "x", "q1": "y"})}
    blind_export.export(runs, qs, tmp_path / "m", seed=1)
    assert [r["question_id"] for r in marks.read_csv(tmp_path / "m" / "key.csv")] == ["q1"]


def test_command_line_round_trip(tmp_path):
    import run_scoring
    qs = full_set(5)
    with open(tmp_path / "q.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=qmod.COLUMNS)
        w.writeheader()
        w.writerows(qs)
    assert run_scoring.main(["validate", str(tmp_path / "q.csv")]) == 0
    runs = {"A": make_run(tmp_path, "ra", {r["id"]: "alpha beta" for r in qs[:3]})}
    assert run_scoring.main(["retrieval", str(tmp_path / "q.csv"), str(tmp_path / "ret.csv"), f"A={runs['A']}"]) == 0
    assert (tmp_path / "ret.csv").read_text().startswith("config_id,question_id,category")
    assert run_scoring.main(["export", str(tmp_path / "q.csv"), str(tmp_path / "m"), "3", f"A={runs['A']}"]) == 0
    assert run_scoring.main(["join", str(tmp_path / "m"), str(tmp_path / "j")]) == 0


def test_summary_tables_show_repeats_and_mark_differences():
    import summary_tables as st
    rows = []
    for rep, correct in (("1", 5), ("2", 5), ("3", 6)):
        for i in range(8):
            rows.append({"config_id": f"A-llama3@{rep}", "category": "self_contained", "correct": "yes" if i < correct else "no"})
    for rep in ("1", "2"):
        for i in range(8):
            rows.append({"config_id": f"B-llama3@{rep}", "category": "self_contained", "correct": "yes" if i < 3 else "no"})
    table = st.correct_table(rows)
    lines = table.splitlines()
    assert lines[0].startswith("| Configuration | self_contained | condition_dependent")
    a = next(l for l in lines if l.startswith("| A-llama3"))
    b = next(l for l in lines if l.startswith("| B-llama3"))
    assert "5 of 8 / 5 of 8 / 6 of 8*" in a and "3 of 8 / 3 of 8 |" in b and "*" not in b
    assert a.count("n/a") == 4


def test_retrieval_table_leaves_out_unanswerable_and_exclusion_table_counts_applicability():
    import summary_tables as st
    rows = [{"config_id": "A@1", "category": "applicability", "any_reference_page": True, "exclusion_retrieved": False},
            {"config_id": "A@1", "category": "applicability", "any_reference_page": True, "exclusion_retrieved": True},
            {"config_id": "A@1", "category": "unanswerable", "any_reference_page": None, "exclusion_retrieved": None}]
    ret = st.retrieval_table(rows).splitlines()[-1]
    exc = st.exclusion_table(rows).splitlines()[-1]
    assert "2 of 2" in ret and ret.count("n/a") == 4
    assert "1 of 2" in exc and exc.count("n/a") == 4


def test_identical_answers_are_shown_once_and_different_ones_apart(tmp_path):
    qs = [question("q0", "self_contained"), question("q1", "self_contained")]
    runs = {"A": make_run(tmp_path, "ra", {"q0": "same text", "q1": "only A"}),
            "B": make_run(tmp_path, "rb", {"q0": "same   text", "q1": "only B"})}
    key = blind_export.export(runs, qs, tmp_path / "m", seed=2)
    blind = marks.read_csv(tmp_path / "m" / "sheet_blind.csv")
    assert len(blind) == 3 and len(key) == 4
    shared = [k["answer_id"] for k in key if k["question_id"] == "q0"]
    assert len(set(shared)) == 1 and {k["config_id"] for k in key if k["answer_id"] == shared[0]} == {"A", "B"}


def test_rubric_scores_each_category():
    import rubric
    q = lambda cat, elems: {"category": cat, "required_elements": elems, "question": "How is the 5 psi limit set?"}
    r = rubric.score_answer(q("self_contained", "close the breaker|set the selector~select RUN"), "First close the breaker, then select RUN.")
    assert r["correct"] == "yes" and r["complete"] == "yes" and r["elements_found"] == "2 of 2"
    r = rubric.score_answer(q("self_contained", "close the breaker|set the selector~select RUN"), "Close the breaker.")
    assert r["correct"] == "partial" and r["complete"] == "no" and r["elements_found"] == "1 of 2"
    r = rubric.score_answer(q("self_contained", "close the breaker"), "Press the green button.")
    assert r["correct"] == "no"
    r = rubric.score_answer(q("unanswerable", "ABSTAIN"), "The manual does not cover that.")
    assert r["correct"] == "yes" and r["abstained"] == "yes"
    r = rubric.score_answer(q("unanswerable", "ABSTAIN"), "It is 40 psi.")
    assert r["correct"] == "no" and r["unsupported_content"] == "yes" and r["numbers_not_in_context"] == "40"
    r = rubric.score_answer(q("applicability", "applies only to~only for MKII"), "This applies only to the new style.")
    assert r["respects_applicability"] == "yes"
    r = rubric.score_answer(q("applicability", "applies only to~only for MKII"), "Step 1. Open the door. Step 2. Press test.")
    assert r["respects_applicability"] == "no" and r["gave_procedure"] == "yes"
    r = rubric.score_answer(q("self_contained", "5 psi"), "Set it to 5 psi.", retrieved_text="limit of 5 psi")
    assert r["unsupported_content"] == "no"


def test_agreement_counts_n_of_n_per_criterion():
    import agreement
    pairs = [({"correct": "yes", "complete": "yes", "abstained": "no", "respects_applicability": "", "unsupported_content": "no"},
              {"correct": "yes", "complete": "no", "abstained": "no", "respects_applicability": "", "unsupported_content": ""}),
             ({"correct": "partial", "complete": "no", "abstained": "no", "respects_applicability": "yes", "unsupported_content": "no"},
              {"correct": "no", "complete": "no", "abstained": "yes", "respects_applicability": "yes", "unsupported_content": "no"})]
    r = agreement.agreement(pairs)
    assert r["correct"] == {"N": 2, "agree": 1, "agree_yes_vs_not_yes": 2}
    assert r["complete"] == {"N": 2, "agree": 1, "agree_yes_vs_not_yes": 1}
    assert r["abstained"]["agree"] == 1 and r["respects_applicability"]["N"] == 1 and r["unsupported_content"]["N"] == 1
    assert "| correct | 1 of 2 | 2 of 2 |" in agreement.render(r)


def test_command_line_scoring_and_agreement(tmp_path):
    import run_scoring
    qs = [question(f"q{i}", "self_contained", elems="alpha|beta") for i in range(4)]
    with open(tmp_path / "q.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=qmod.COLUMNS)
        w.writeheader()
        w.writerows(qs)
    answers = {"q0": "alpha and beta", "q1": "alpha only", "q2": "I do not know", "q3": "alpha and beta again"}
    run = make_run(tmp_path, "ra", answers)
    assert run_scoring.main(["score", str(tmp_path / "q.csv"), str(tmp_path / "auto.csv"), f"A@1={run}"]) == 0
    auto = marks.read_csv(tmp_path / "auto.csv")
    assert [r["correct"] for r in auto] == ["yes", "partial", "no", "yes"]
    assert run_scoring.main(["export", str(tmp_path / "q.csv"), str(tmp_path / "m"), "5", f"A@1={run}"]) == 0
    blind = marks.read_csv(tmp_path / "m" / "sheet_blind.csv")
    for r in blind:   # the person agrees with the rubric except on the partial answer
        r["correct"] = "no" if r["answer"] == "alpha only" else ("yes" if "beta" in r["answer"] else "no")
        r["complete"] = "yes" if "beta" in r["answer"] else "no"
        r["abstained"] = "yes" if "know" in r["answer"] else "no"
    with open(tmp_path / "m" / "sheet_blind.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(blind[0]))
        w.writeheader()
        w.writerows(blind)
    assert run_scoring.main(["agreement", str(tmp_path / "m")]) == 0


def test_verification_pack_has_one_row_per_question_and_empty_verdict_columns(tmp_path):
    import make_verification_pack as mv
    from openpyxl import load_workbook
    draft = tmp_path / "draft.csv"
    with open(draft, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=mv.COLUMNS)
        w.writeheader()
        for i in range(3):
            w.writerow({"id": f"q{i}", "set": "main", "forbidden_elements": "x", "category": "self_contained", "question": f"question {i}", "reference_answer": "ref",
                        "source_pages": "4", "exclusion_pages": "", "required_elements": "a|b", "passage": f"passage text {i}"})
    assert mv.main(str(draft), str(tmp_path / "pack")) == 3
    wb = load_workbook(tmp_path / "pack" / "verification_pack.xlsx")
    ws = wb["questions"]
    header = [c.value for c in ws[1]]
    assert header[-4:] == ["verdict", "correction", "verified_by", "verified_on"]
    assert ws.max_row == 4 and ws.cell(row=2, column=header.index("passage") + 1).value == "passage text 0"
    assert all(ws.cell(row=r, column=header.index("verdict") + 1).value in (None, "") for r in range(2, 5))
    rows = marks.read_csv(tmp_path / "pack" / "verification_pack.csv")
    assert len(rows) == 3 and rows[0]["verdict"] == ""


def test_wilson_mcnemar_and_holm():
    import analysis
    lo, hi = analysis.wilson(8, 10)
    assert round(lo, 3) == 0.490 and round(hi, 3) == 0.943
    assert analysis.wilson(0, 0) == (None, None)
    assert analysis.mcnemar_exact(0, 0) == 1.0
    assert abs(analysis.mcnemar_exact(0, 5) - 0.0625) < 1e-12
    assert abs(analysis.mcnemar_exact(2, 8) - 0.109375) < 1e-12
    assert analysis.mcnemar_exact(5, 5) == 1.0
    assert analysis.holm([0.01, 0.04, 0.03]) == [0.03, 0.06, 0.06]


def test_comparisons_pair_questions_and_count_discordant(tmp_path):
    import analysis, json
    rows = []
    for q, (a, b) in enumerate([("yes", "no")] * 6 + [("no", "yes")] + [("yes", "yes")] * 3 + [("partial", "no")]):
        rows += [{"config_id": "A-x@1", "question_id": f"q{q}", "outcome": a}, {"config_id": "B-x@1", "question_id": f"q{q}", "outcome": b}]
    res = analysis.run_comparisons(rows, [{"id": "P1", "a": "A-x", "b": "B-x", "rep": 1}])[0]
    assert (res["N_paired"], res["a_correct"], res["b_correct"], res["a_only"], res["b_only"]) == (11, 9, 4, 6, 1)
    assert abs(res["p_exact"] - 0.125) < 1e-12 and res["p_holm"] == res["p_exact"]
    partial = analysis.run_comparisons(rows, [{"id": "P1", "a": "A-x", "b": "B-x"}], partial_credit=True)[0]
    assert partial["a_only"] == 7
    csvp = tmp_path / "o.csv"
    csvp.write_text("config_id,question_id,outcome\n" + "".join(f'{r["config_id"]},{r["question_id"]},{r["outcome"]}\n' for r in rows))
    cmp = tmp_path / "c.json"
    cmp.write_text(json.dumps({"comparisons": [{"id": "P1", "a": "A-x", "b": "B-x", "rep": 1}]}))
    analysis.run(str(csvp), str(cmp), str(tmp_path / "out"))
    assert (tmp_path / "out_counts.csv").exists() and (tmp_path / "out_comparisons_with_partial.csv").exists()


def test_kappa_and_two_markers():
    import agreement
    k, n, agree = agreement.cohen_kappa([("yes", "yes")] * 20 + [("no", "no")] * 15 + [("yes", "no")] * 5 + [("no", "yes")] * 10)
    assert (n, agree) == (50, 35) and abs(k - 0.4) < 1e-9
    assert agreement.cohen_kappa([("yes", "yes")] * 4)[0] is None
    assert agreement.cohen_kappa([("", "yes")]) == (None, 0, 0)
    res = agreement.two_markers({"a1": {"correct": "yes"}, "a2": {"correct": "no"}}, {"a1": {"correct": "yes"}, "a2": {"correct": "yes"}, "a3": {"correct": "no"}})
    assert res["correct"]["N"] == 2 and res["correct"]["agree"] == 1
    assert "1 of 2" in agreement.render_two_markers(res)


def test_second_marker_sample_is_blind_and_reproducible(tmp_path):
    import blind_export, csv
    rows = [{"answer_id": f"ans-{i:04d}", "question": "q", "reference_answer": "r", "required_elements": "e", "category": "c", "answer": "a",
             "correct": "yes", "complete": "yes", "respects_applicability": "", "abstained": ""} for i in range(1, 101)]
    with open(tmp_path / "sheet_blind.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    ids = blind_export.second_marker_sample(tmp_path, 60, 7)
    assert len(ids) == 60 and len(set(ids)) == 60 and ids == blind_export.second_marker_sample(tmp_path, 60, 7)
    sample = list(csv.DictReader(open(tmp_path / "sheet_second_marker.csv")))
    assert all(r["correct"] == "" and r["complete"] == "" for r in sample) and "config_id" not in sample[0]


def test_context_texts_use_the_given_pages_in_the_oracle_bracket():
    import blind_export
    assert blind_export.context_texts({"retrieved": [], "given_text": "page text"}, {}) == ["page text"]
    assert blind_export.context_texts({"retrieved": [], "given_text": ""}, {}) == []
    assert blind_export.context_texts({"retrieved": [{"chunk_id": "c1"}]}, {"c1": {"text": "t"}}) == ["t"]


def test_summary_tables_name_run_and_commit(tmp_path):
    import check_index
    (tmp_path / "index.md").write_text("| `e1` | table |\n")
    (tmp_path / "summary.md").write_text("## A\n\nSource: run 1, commit abc. Index entry: `e1`.\n\n| a | b |\n|---|---|\n")
    assert check_index.problems(tmp_path) == []
    (tmp_path / "summary.md").write_text("## A\n\n| a | b |\n|---|---|\n")
    assert len(check_index.problems(tmp_path)) == 1
    (tmp_path / "summary.md").write_text("## A\n\nSource: run 1, commit abc. Index entry: `zz`.\n\n| a |\n")
    assert "zz" in check_index.problems(tmp_path)[0]


def test_freeze_manifest_detects_a_changed_file(tmp_path, monkeypatch):
    import freeze_manifest as fm
    (tmp_path / "a.txt").write_text("one")
    monkeypatch.setattr(fm, "ROOT", tmp_path)
    monkeypatch.setattr(fm, "FROZEN", ["a.txt"])
    fm.write(tmp_path / "m")
    assert fm.check(tmp_path / "m") == []
    (tmp_path / "a.txt").write_text("two")
    assert fm.check(tmp_path / "m") == ["a.txt: changed"]
    (tmp_path / "a.txt").unlink()
    assert fm.check(tmp_path / "m") == ["a.txt: missing"]


def test_forbidden_elements_fail_an_answer_and_the_order_is_reported():
    import rubric
    q = {"category": "applicability", "question": "q", "required_elements": "345 kPa~50 psi", "forbidden_elements": "276 kPa~40 psi"}
    assert rubric.score_answer(q, "The setting is 345 kPa.")["correct"] == "yes"
    out = rubric.score_answer(q, "345 kPa (50 psi), where other models use 40 psi.")
    assert out["correct"] == "no" and out["complete"] == "no" and out["respects_applicability"] == "no" and out["forbidden_found"] == "276 kPa~40 psi"
    assert rubric.score_answer({**q, "forbidden_elements": "yes"}, "Their eyes are on 345 kPa.")["correct"] == "yes"  # whole words only
    proc = {"category": "self_contained", "question": "q", "required_elements": "open the valve|press start|close the valve"}
    assert rubric.score_answer(proc, "Open the valve, press START, then close the valve.")["elements_in_order"] == "yes"
    assert rubric.score_answer(proc, "Close the valve. Press start. Open the valve.")["elements_in_order"] == "no"


def test_validate_requires_forbidden_elements_on_applicability_rows_and_can_skip_minimums():
    rows = [question("ap-1", "applicability", exc="4"), question("sc-1", "self_contained")]
    rows[0]["forbidden_elements"] = ""
    errors, _ = qmod.validate(rows, check_counts=False)
    assert errors == ["ap-1: applicability question has no forbidden_elements"]
    rows[0]["forbidden_elements"] = "other value"
    assert qmod.validate(rows, check_counts=False)[0] == []
    assert qmod.validate(rows)[0]  # the category minimums still apply by default
    old = [{k: v for k, v in r.items() if k != "forbidden_elements"} for r in rows]
    errors, warnings = qmod.validate(old, check_counts=False)
    assert errors == [] and any("forbidden_elements" in w for w in warnings)


def test_merge_applies_changed_cells_and_verdicts(tmp_path):
    import merge_verification as mv
    base = {"question": "q", "reference_answer": "r", "source_pages": "1", "exclusion_pages": "", "required_elements": "a", "forbidden_elements": "", "verified_by": "", "verified_on": ""}
    main = [{**base, "id": f"sc-{i}", "category": "self_contained"} for i in range(1, 6)]
    main[0]["forbidden_elements"] = "repo value"  # corrected in the repo after delivery
    delivered = [{**base, "id": r["id"], "set": "main", "verdict": "", "correction": ""} for r in main]
    delivered[0]["forbidden_elements"] = "old value"
    returned = [dict(d) for d in delivered]
    returned[0]["verdict"] = "ok"
    returned[1].update(verdict="fix", reference_answer="r2")
    returned[2].update(verdict="fix", correction="")
    returned[3].update(verdict="drop")
    unread = []
    report = mv.merge(delivered, returned, main, unread, "M", "2026-10-06")
    by_id = {q["id"]: q for q in main}
    assert by_id["sc-1"]["forbidden_elements"] == "repo value" and by_id["sc-1"]["verified_by"] == "M"
    assert by_id["sc-2"]["reference_answer"] == "r2" and by_id["sc-2"]["verified_by"] == "M"
    assert by_id["sc-3"]["verified_by"] == "" and report["unresolved"][0].startswith("sc-3")
    assert by_id["sc-4"].get("_drop") and report["dropped"] == ["sc-4"] and report["unverified"] == ["sc-5"]


def test_parse_correction_splits_question_answer_pages_and_notes():
    import merge_verification as mv
    fix = mv.parse_correction("Question: New question?\nAnswer: First step. Second step. PDF pages 78-79. Note: page 101 differs.")
    assert fix == {"question": "New question?", "answer": "First step. Second step.", "pages": "78;79", "notes": "Note: page 101 differs."}
    assert mv.parse_correction("Answer: One. PDF page 42. Added a check.")["pages"] == "42"
    unanswerable = mv.parse_correction("Answer: The manual does not say. Do not guess. The original claim was wrong: PDF page 111 also covers it. The question remains unanswerable.")
    assert unanswerable["answer"] == "The manual does not say. Do not guess." and unanswerable["pages"] is None
    assert mv.parse_correction("only a note")["answer"] is None


def test_short_alternatives_match_whole_words_and_longer_ones_match_substrings():
    import rubric
    n = rubric.normalise("Top up to H. Turn the controller OFF, then offer help. Ventilation first.")
    assert rubric.element_met(n, "H mark~H") and rubric.element_met(n, "OFF") and rubric.element_met(n, "ventilat")
    assert not rubric.element_met(rubric.normalise("the heater is offered"), "H") and not rubric.element_met(rubric.normalise("offered"), "OFF")
    assert not rubric.element_met(rubric.normalise("TB-10 is the contactor"), "TB-1") and rubric.element_met(rubric.normalise("check TB-1 first"), "TB-1")


def test_self_check_flags_an_answer_that_misses_its_own_element():
    row = {"id": "x", "category": "self_contained", "question": "q", "reference_answer": "Top up to H.", "source_pages": "1", "exclusion_pages": "",
           "required_elements": "H mark~high mark", "forbidden_elements": "", "verified_by": "SA", "verified_on": "2026-10-05"}
    assert qmod.self_check([row])[0].startswith("x: required elements not met")
    assert qmod.self_check([{**row, "required_elements": "H mark~high mark~H"}]) == []
    ap = {**row, "id": "y", "category": "applicability", "required_elements": "H", "forbidden_elements": "top up"}
    assert "trips forbidden" in qmod.self_check([ap])[0]
    un = {**row, "id": "z", "category": "unanswerable", "reference_answer": "The manual gives no value.", "required_elements": "ABSTAIN"}
    assert qmod.self_check([un]) == []
    assert qmod.self_check([{**un, "reference_answer": "It is 5."}])[0].startswith("z: the reference answer does not read as an abstention")
