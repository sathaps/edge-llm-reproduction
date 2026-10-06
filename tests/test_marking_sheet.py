import csv, json, random, sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import marking_sheet as ms  # noqa: E402

CELLS = ["A-tinyllama", "B-tinyllama", "A-llama3", "B-llama3", "A-mistral", "B-mistral"]


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    base = tmp_path_factory.mktemp("e1")
    qs = list(csv.DictReader(open(ROOT / "questions" / "questions.csv")))
    rng = random.Random(1)
    for cell in CELLS:
        d = base / "e1" / cell / "rep-1"
        d.mkdir(parents=True)
        with open(d / "answers.jsonl", "w") as f:
            for q in qs:
                k = rng.choice(["same", "spaced", "empty", "own"])
                ans = {"same": "The manual gives no such value.", "spaced": "  The manual   gives no\nsuch value. ", "empty": "",
                       "own": f"Line one\nline two from {cell} for {q['id']}"}[k]
                f.write(json.dumps({"question_id": q["id"], "answer": ans, "run_id": 99}) + "\n")
    out = base / "out"
    meta = ms.build(base / "e1", 1, out, 4242, 7)
    return qs, out, meta


def rows(path, tab=ms.TAB1):
    return ms.read_tab(load_workbook(path)[tab])[1]


def fill(src, dst, value=None):
    wb = load_workbook(src)
    ws = wb[ms.TAB1]
    head = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        cat = ws.cell(r, head.index("category") + 1).value
        for name, ok in ms.MARKS.items():
            if name == "respects_applicability" and cat != "applicability":
                continue
            ws.cell(r, head.index(name) + 1).value = ok[0] if value is None else value
    wb.save(dst)


def test_merge_and_counts(built):
    qs, out, meta = built
    sheet = rows(out / "e1_rep1_marking_sheet.xlsx")
    assert meta["answers_total"] == 240 and meta["distinct_answers"] == len(sheet) < 240
    assert sum(meta["distinct_by_category"].values()) == len(sheet)
    texts = [(r["question_id"], ms.norm(r["answer"])) for r in sheet]
    assert len(texts) == len(set(texts))
    assert any(r["answer"] == ms.EMPTY for r in sheet)
    assert any("\n" in r["answer"] for r in sheet)


def test_question_order_and_ids(built):
    qs, out, _ = built
    sheet = rows(out / "e1_rep1_marking_sheet.xlsx")
    order = [q["id"] for q in qs]
    seen = []
    for r in sheet:
        if not seen or seen[-1] != r["question_id"]:
            seen.append(r["question_id"])
    assert seen == order
    ids = [r["answer_id"] for r in sheet]
    assert len(set(ids)) == len(ids)


def test_blind(built):
    _, out, _ = built
    words = [c.lower() for c in CELLS] + ["tinyllama", "llama3", "mistral", "run_id", "retriev", "chunk", "score", "rep-1"]
    for name in ("e1_rep1_marking_sheet.xlsx", "e1_rep1_second_marker_sample.xlsx"):
        wb = load_workbook(out / name)
        for ws in wb:
            head = [c.value for c in ws[1]]
            assert head == ms.HEADER
            for row in ws.iter_rows(min_row=2, values_only=True):
                cells = dict(zip(head, row))
                blob = " ".join(str(v) for k, v in cells.items() if k != "answer" and v is not None).lower()
                assert not [w for w in words if w in blob], (name, ws.title, blob[:80])


def test_header_dropdowns_and_freeze(built):
    _, out, _ = built
    wb = load_workbook(out / "e1_rep1_marking_sheet.xlsx")
    ws = wb[ms.TAB1]
    assert [c.value for c in ws[1]] == ms.HEADER and ws.freeze_panes == "B2"
    formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"yes,partial,no"' in formulas and '"yes,no"' in formulas
    assert wb.sheetnames == [ms.TAB1]


def test_sample(built):
    _, out, meta = built
    sample = rows(out / "e1_rep1_second_marker_sample.xlsx")
    sheet = {r["answer_id"] for r in rows(out / "e1_rep1_marking_sheet.xlsx")}
    assert len(sample) == 60 and {r["answer_id"] for r in sample} <= sheet
    assert set(meta["sample_by_category"].values()) == {12}
    assert set(meta["sample_by_stratum_cell"].values()) == {10}
    assert all(not r["correct"] for r in sample)


def test_check_accepts_filled_and_names_problems(built, tmp_path):
    _, out, _ = built
    orig = out / "e1_rep1_marking_sheet.xlsx"
    good = tmp_path / "good.xlsx"
    fill(orig, good)
    assert ms.check(good, orig, tmp_path / "marks.csv") == []
    assert len(list(csv.DictReader(open(tmp_path / "marks.csv")))) == len(rows(orig))
    assert ms.check(orig, orig)[0].startswith(f"{ms.TAB1} row 2")
    bad = tmp_path / "bad.xlsx"
    fill(orig, bad, value="maybe")
    problems = ms.check(bad, orig)
    assert problems and "allowed" in problems[0]
    wb = load_workbook(good)
    ws = wb[ms.TAB1]
    ws.cell(3, ms.HEADER.index("answer") + 1).value = "edited"
    ws.cell(4, ms.HEADER.index("correct") + 1).value = None
    wb.save(tmp_path / "edit.xlsx")
    problems = ms.check(tmp_path / "edit.xlsx", orig)
    assert any("row 3" in p and "answer was changed" in p for p in problems)
    assert any("row 4" in p and "correct is empty" in p for p in problems)
